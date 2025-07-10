from openpyxl import Workbook
from openpyxl import load_workbook
import datetime
from chart import Chart
from openpyxl.chart import Series, Reference, BarChart
from openpyxl.chart.label import DataLabelList

class MonthlyReportChart(Chart):
    
    def __init__(self,chart_template_file,chart_file):
        super().__init__(chart_template_file,chart_file)
        
    def update_worksheet(self,stats,name,start_row,col_names):
        today = datetime.date.today()
        first = today.replace(day=1)
        previous_month = first - datetime.timedelta(days=1)
        month_lookup = previous_month.strftime("%Y-%m")
  
        srs=self.get_worksheet(name)
        srs["B1"]=today
        srs["D1"]=month_lookup
        i=start_row
        for stat in stats:
            cols=[srs[f"A{i}"],srs[f"B{i}"],srs[f"C{i}"],srs[f"D{i}"],srs[f"E{i}"]]
            i=i+1
            for j in range(0,len(cols)):
                cols[j].value=stat[col_names[j]] 
    
    def get_chart_height(self,name="Report"):
       srs=self.get_worksheet(name)
       # 0.035 is conversion ratio from point to cm
       height=srs.row_dimensions[1].height*0.035
       return height

    def update_monthly_views(self,stats,coverage_month,report_date,col_names):
        ws_report=self.get_worksheet("Report")
        ws_report["I1"].value=str(coverage_month)
        ws_report["C4"].value=str(report_date)

        ws_stat=self.get_worksheet("Statistics")
        row=3
        for stat in stats:
            cols=[ws_stat[f"A{row}"],ws_stat[f"B{row}"],ws_stat[f"C{row}"],ws_stat[f"D{row}"],ws_stat[f"E{row}"]]
            row=row+1
            for j in range(0,len(cols)):
               cols[j].value=stat[col_names[j]]
 

    def add_chart(self,x_title="",y_title="",min_data_col=2,max_data_col=2,minrow=3,maxrow=10,loc="A10", height=15):
       chart1 = BarChart()
       chart1.type = "bar"
       chart1.style = 11
       chart1.width = 25
       chart1.height = height
       chart1.x_axis.title = x_title 
       chart1.y_axis.title = y_title 

       ws1=self.get_worksheet("Statistics")
       data = Reference(ws1, min_col=min_data_col, min_row=minrow, max_row=maxrow, max_col=max_data_col)
       cats = Reference(ws1, min_col=1, min_row=minrow, max_row=maxrow)
       chart1.add_data(data, titles_from_data=False)
       chart1.set_categories(cats)
       chart1.shape = 4
       chart1.dataLabels = DataLabelList() 
       chart1.dataLabels.showVal = True
       chart1.legend=None
       ws=self.get_worksheet("Report")
       ws.add_chart(chart1, loc)

